"""
tests/test_output_format.py
----------------------------
Unit tests for write_excel_output format validation.

Tests verify:
- Correct columns in output
- Effort/Timeline label expansion (EN/ES/PT)
- Category validation and correction
- Blocked URL filtering
- Impact column comment
- Corporate styling basics

Run: python -m pytest tests/test_output_format.py -v
"""
import sys
import pathlib
import tempfile
import os

import pytest
import openpyxl

# Add project root to path
sys.path.insert(0, str(pathlib.Path(__file__).parent.parent))

# Mock environment variables before importing modules that need them
os.environ.setdefault("HANA_HOST", "localhost")
os.environ.setdefault("HANA_PORT", "443")
os.environ.setdefault("HANA_USER", "test")
os.environ.setdefault("HANA_PASSWORD", "test")
os.environ.setdefault("HANA_SCHEMA", "SVA2")
os.environ.setdefault("AICORE_AUTH_URL", "http://localhost")
os.environ.setdefault("AICORE_CLIENT_ID", "test")
os.environ.setdefault("AICORE_CLIENT_SECRET", "test")
os.environ.setdefault("AICORE_BASE_URL", "http://localhost")
os.environ.setdefault("EMBEDDING_DEPLOYMENT_ID", "test")
os.environ.setdefault("EMBEDDING_MODEL_NAME", "test")


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------
@pytest.fixture
def sample_input_xlsx(tmp_path):
    """Create a minimal input Excel for write_excel_output to derive filename from."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Pain Point", "Solution"])
    ws.append(["Test pain point", "Ariba Sourcing"])
    path = tmp_path / "test_input.xlsx"
    wb.save(str(path))
    return str(path)


@pytest.fixture
def sample_rows():
    """Standard test rows with various field values."""
    return [
        {
            "idx": 0,
            "pain_point": "No approval workflow configured",
            "solution": "Ariba Sourcing",
            "Recommendations": "Configure approval workflows for sourcing events.",
            "Category": "Feature Adoption",
            "Effort": "Low",
            "Timeline": "Quick Win",
            "Benefits": "Faster approvals and compliance.",
            "Documentation": [
                {"title": "Approval Workflows", "url": "https://help.sap.com/docs/abc123def456abc123def456abc123de/guide"}
            ],
            "Impact": "",
            "Ariba Next-Gen": "No feature identified.",
            "Value KPIs": "No KPI data available.",
        },
        {
            "idx": 1,
            "pain_point": "La integración con SAP no funciona correctamente",
            "solution": "Ariba Buying",
            "Recommendations": "Revisar la configuración de integración CIG.",
            "Category": "innovation",  # lowercase — should be corrected
            "Effort": "medio",  # Spanish — should expand
            "Timeline": "corto plazo",  # Spanish — should expand
            "Benefits": "Mejora en la sincronización de datos.",
            "Documentation": [
                {"title": "Blocked", "url": "https://learning.sap.com/course/123"}
            ],
            "Impact": "",
            "Ariba Next-Gen": "",
            "Value KPIs": "",
        },
    ]


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------
class TestWriteExcelOutput:
    """Tests for write_excel_output function."""

    def _write(self, input_path, rows, tmp_path):
        """Helper to call write_excel_output with a temp output path."""
        from server.recommend import write_excel_output
        output = str(tmp_path / "output.xlsx")
        result_path, rows_written = write_excel_output(input_path, rows, output)
        return result_path, rows_written

    def test_correct_columns(self, sample_input_xlsx, sample_rows, tmp_path):
        """Output should have all expected columns in correct order."""
        path, _ = self._write(sample_input_xlsx, sample_rows, tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert "Pain Point" in headers
        assert "Solution" in headers
        assert "Recommendations" in headers
        assert "Category" in headers
        assert "Effort" in headers
        assert "Timeline" in headers
        assert "Benefits" in headers
        assert "Documentation" in headers
        assert "Impact" in headers
        assert "Ariba Next-Gen" in headers
        assert "Value KPIs" in headers

    def test_effort_expansion_english(self, sample_input_xlsx, sample_rows, tmp_path):
        """English short labels should expand to full descriptive strings."""
        path, _ = self._write(sample_input_xlsx, sample_rows, tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        # Find Effort column
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        effort_col = headers.index("Effort") + 1
        # Row 2 (idx=0) had "Low"
        assert ws.cell(row=2, column=effort_col).value == "Low (1 – 3 Days)"

    def test_effort_expansion_spanish(self, sample_input_xlsx, sample_rows, tmp_path):
        """Spanish Effort aliases should expand correctly."""
        path, _ = self._write(sample_input_xlsx, sample_rows, tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        effort_col = headers.index("Effort") + 1
        # Row 3 (idx=1) had "medio"
        assert ws.cell(row=3, column=effort_col).value == "Medium (1 – 3 Weeks)"

    def test_timeline_expansion_spanish(self, sample_input_xlsx, sample_rows, tmp_path):
        """Spanish Timeline aliases should expand correctly."""
        path, _ = self._write(sample_input_xlsx, sample_rows, tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        timeline_col = headers.index("Timeline") + 1
        # Row 3 (idx=1) had "corto plazo"
        assert ws.cell(row=3, column=timeline_col).value == "Short Term (1 – 3 Weeks)"

    def test_category_correction(self, sample_input_xlsx, sample_rows, tmp_path):
        """Lowercase category variants should be corrected to canonical form."""
        path, _ = self._write(sample_input_xlsx, sample_rows, tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        cat_col = headers.index("Category") + 1
        # Row 3 (idx=1) had "innovation" (lowercase)
        assert ws.cell(row=3, column=cat_col).value == "Innovation"

    def test_blocked_urls_filtered(self, sample_input_xlsx, sample_rows, tmp_path):
        """learning.sap.com URLs should be stripped from output."""
        path, _ = self._write(sample_input_xlsx, sample_rows, tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        doc_col = headers.index("Documentation") + 1
        # Row 3 (idx=1) had only a learning.sap.com URL — should be empty
        doc_value = ws.cell(row=3, column=doc_col).value
        assert not doc_value or "learning.sap.com" not in doc_value

    def test_impact_header_comment(self, sample_input_xlsx, sample_rows, tmp_path):
        """Impact header cell should have an explanatory comment."""
        path, _ = self._write(sample_input_xlsx, sample_rows, tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        impact_col = headers.index("Impact") + 1
        cell = ws.cell(row=1, column=impact_col)
        assert cell.comment is not None
        assert "consultant" in cell.comment.text.lower()

    def test_rows_written_count(self, sample_input_xlsx, sample_rows, tmp_path):
        """Should report correct number of rows written."""
        _, rows_written = self._write(sample_input_xlsx, sample_rows, tmp_path)
        assert rows_written == 2

    def test_corporate_header_style(self, sample_input_xlsx, sample_rows, tmp_path):
        """Header row should have corporate navy blue fill."""
        path, _ = self._write(sample_input_xlsx, sample_rows, tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        header_cell = ws.cell(row=1, column=1)
        assert header_cell.font.bold is True
        assert header_cell.font.color.rgb == "00FFFFFF"  # white text
        assert header_cell.fill.fgColor.rgb == "0000144A"  # navy blue

    def test_portuguese_effort_expansion(self, sample_input_xlsx, tmp_path):
        """Portuguese Effort aliases should expand correctly."""
        rows = [{
            "idx": 0,
            "pain_point": "Test",
            "solution": "Ariba Sourcing",
            "Effort": "baixo",
            "Timeline": "curto prazo",
        }]
        path, _ = self._write(sample_input_xlsx, rows, tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        effort_col = headers.index("Effort") + 1
        timeline_col = headers.index("Timeline") + 1
        assert ws.cell(row=2, column=effort_col).value == "Low (1 – 3 Days)"
        assert ws.cell(row=2, column=timeline_col).value == "Short Term (1 – 3 Weeks)"

    def test_valid_url_passes_through(self, sample_input_xlsx, tmp_path):
        """Valid help.sap.com URLs with GUID should pass through."""
        rows = [{
            "idx": 0,
            "pain_point": "Test",
            "solution": "Ariba Sourcing",
            "Documentation": [
                {"title": "Guide", "url": "https://help.sap.com/docs/abc123def456abc123def456abc123de/topic/detail"}
            ],
        }]
        path, _ = self._write(sample_input_xlsx, rows, tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        doc_col = headers.index("Documentation") + 1
        doc_value = ws.cell(row=2, column=doc_col).value
        assert "help.sap.com" in doc_value
